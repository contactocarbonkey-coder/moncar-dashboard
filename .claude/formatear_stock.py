"""Aplica formato profesional y validaciones al STOCK_MONCAR.xlsx"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

PATH = r'C:\Users\franc\OneDrive\Documentos\KeyCarbon\claude MONCAR RMH\STOCK_MONCAR.xlsx'

# ── Paleta de colores ─────────────────────────────────────────────────────────
NAVY   = "1F3864"   # azul oscuro — cabeceras y títulos
BLUE   = "2E75B6"   # azul medio — fila de totales
BL_LT  = "DEEAF1"   # azul muy claro — filas pares
GRAY_A = "EFEFEF"   # gris claro — campos automáticos
WHITE  = "FFFFFF"
TXT_WH = "FFFFFF"
TXT_DK = "1F1F1F"
TXT_GR = "666666"

# ── Helpers ───────────────────────────────────────────────────────────────────
def fill(h):
    return PatternFill("solid", fgColor=h)

def fnt(color=TXT_DK, bold=False, size=11, italic=False):
    return Font(color=color, bold=bold, size=size, italic=italic, name="Calibri")

def aln(h="center", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

_s = Side(style="thin", color="C0C0C0")
BRD = Border(left=_s, right=_s, top=_s, bottom=_s)

def style_header(cell):
    cell.fill = fill(NAVY)
    cell.font = fnt(TXT_WH, bold=True, size=10)
    cell.alignment = aln("center", "center", wrap=True)
    cell.border = BRD

def style_total(cell):
    cell.fill = fill(BLUE)
    cell.font = fnt(TXT_WH, bold=True, size=10)
    cell.alignment = aln("center", "center")
    cell.border = BRD

wb = openpyxl.load_workbook(PATH)

# ═══════════════════════════════════════════════════════════════════════════════
# HOJA: Productos
# ═══════════════════════════════════════════════════════════════════════════════
ws = wb["Productos"]
ws.sheet_view.showGridLines = False

ws.merge_cells("A1:G1")
ws["A1"].value = "PRODUCTOS MONCAR  —  Stock base inicial"
ws["A1"].fill = fill(NAVY)
ws["A1"].font = fnt(TXT_WH, bold=True, size=13)
ws["A1"].alignment = aln("left", "center")
ws.row_dimensions[1].height = 34

ws["A2"] = "TOTALES"
ws.row_dimensions[2].height = 22
for c in range(1, 8):
    style_total(ws.cell(row=2, column=c))
    if c >= 4:
        ws.cell(row=2, column=c).number_format = "#,##0"

ws.row_dimensions[3].height = 40
for c in range(1, 8):
    style_header(ws.cell(row=3, column=c))

for r in range(4, ws.max_row + 1):
    rf = fill(BL_LT) if r % 2 == 0 else fill(WHITE)
    for c in range(1, 8):
        cell = ws.cell(row=r, column=c)
        cell.fill = rf
        cell.border = BRD
        cell.font = fnt(size=10)
        cell.alignment = aln("left", "center") if c <= 2 else aln("center", "center")
        if c >= 3:
            cell.number_format = "#,##0"

for col, w in [("A",11),("B",22),("C",13),("D",13),("E",15),("F",13),("G",12)]:
    ws.column_dimensions[col].width = w
ws.freeze_panes = "C4"

# ═══════════════════════════════════════════════════════════════════════════════
# HOJA: Movimientos
# ═══════════════════════════════════════════════════════════════════════════════
ws = wb["Movimientos"]
ws.sheet_view.showGridLines = False

TIPOS = [
    "Ingreso",
    "Embolsado",
    "Embolsado Premium",
    "Salida Venta Franco",
    "Salida Lita",
    "Entrega a Juani",
    "Salida Venta Juani",
    "Préstamo dado",
    "Préstamo recibido",
    "Recupero préstamo dado",
    "Devolución simple",
    "Devolución préstamo recibido",
]
ESTADOS = ["No embolsado", "Embolsado", "Premium"]
QUIEN   = ["Franco", "Sofi", "Bauti"]

ws.merge_cells("A1:I1")
ws["A1"].value = "MOVIMIENTOS DE STOCK MONCAR  —  Registrá un movimiento por fila"
ws["A1"].fill = fill(NAVY)
ws["A1"].font = fnt(TXT_WH, bold=True, size=13)
ws["A1"].alignment = aln("left", "center")
ws.row_dimensions[1].height = 34

ws.merge_cells("A2:I2")
ws["A2"].value = (
    "  Columnas grises (Nº y Año-Mes) son automáticas  •  "
    "Tipo, Estado y Quién embolsó tienen desplegable  •  "
    "Completar: Fecha | Tipo | Estado | COD | Cantidad | Quién embolsó | Comentarios"
)
ws["A2"].fill = fill(BL_LT)
ws["A2"].font = fnt(TXT_DK, size=9, italic=True)
ws["A2"].alignment = aln("left", "center")
ws.row_dimensions[2].height = 18

ws.row_dimensions[3].height = 40
for c in range(1, 10):
    style_header(ws.cell(row=3, column=c))

AUTO_COLS = {1, 9}
for r in range(4, 1004):
    rf = fill(BL_LT) if r % 2 == 0 else fill(WHITE)
    ws.row_dimensions[r].height = 18
    for c in range(1, 10):
        cell = ws.cell(row=r, column=c)
        if c in AUTO_COLS:
            cell.fill = fill(GRAY_A)
            cell.font = fnt(TXT_GR, size=10)
        else:
            cell.fill = rf
            cell.font = fnt(size=10)
        cell.border = BRD
        if c == 1:
            cell.alignment = aln("center", "center")
        elif c == 2:
            cell.alignment = aln("center", "center")
            cell.number_format = "DD/MM/YYYY"
        elif c == 6:
            cell.alignment = aln("center", "center")
            cell.number_format = "#,##0"
        elif c == 9:
            cell.alignment = aln("center", "center")
        else:
            cell.alignment = aln("left", "center")

# Validaciones
def add_dv(ws, formula, sqref, ptitle, prompt, etitle, error):
    dv = DataValidation(type="list", formula1=formula,
                        allow_blank=True, showDropDown=False)
    dv.promptTitle = ptitle
    dv.prompt = prompt
    dv.errorTitle = etitle
    dv.error = error
    dv.showInputMessage = False
    dv.showErrorMessage = True
    dv.sqref = sqref
    ws.add_data_validation(dv)

tipo_formula = '"' + ",".join(TIPOS) + '"'
add_dv(ws, tipo_formula, "C4:C1003",
       "Tipo de movimiento", "Elegi el tipo de movimiento de la lista",
       "Tipo no valido", "Solo se permiten los tipos predefinidos.")

estado_formula = '"' + ",".join(ESTADOS) + '"'
add_dv(ws, estado_formula, "D4:D1003",
       "Estado", "No embolsado / Embolsado / Premium",
       "Estado no valido", "Solo se permiten: No embolsado, Embolsado, Premium.")

quien_formula = '"' + ",".join(QUIEN) + '"'
add_dv(ws, quien_formula, "G4:G1003",
       "Quien embolso", "Elegi quien realizo el embolsado",
       "Valor no valido", "Solo se permiten: Franco, Sofi, Bauti.")

for col, w in [("A",6),("B",13),("C",26),("D",16),("E",14),
               ("F",12),("G",15),("H",30),("I",12)]:
    ws.column_dimensions[col].width = w
ws.freeze_panes = "B4"

# ═══════════════════════════════════════════════════════════════════════════════
# HOJA: Stock
# ═══════════════════════════════════════════════════════════════════════════════
ws = wb["Stock"]
ws.sheet_view.showGridLines = False

ws.merge_cells("A1:J1")
ws["A1"].value = "STOCK MONCAR  —  Se actualiza automáticamente con cada movimiento"
ws["A1"].fill = fill(NAVY)
ws["A1"].font = fnt(TXT_WH, bold=True, size=13)
ws["A1"].alignment = aln("left", "center")
ws.row_dimensions[1].height = 34

ws["A2"] = "TOTALES"
ws.row_dimensions[2].height = 22
for c in range(1, 11):
    style_total(ws.cell(row=2, column=c))
    if c >= 3:
        ws.cell(row=2, column=c).number_format = "#,##0"

ws.row_dimensions[3].height = 40
for c in range(1, 11):
    style_header(ws.cell(row=3, column=c))

VCOL = {7, 8}    # VENDIBLE y TOTAL GENERAL — verde
PCOL = {9, 10}   # PREST y DEBO — dorado

for r in range(4, ws.max_row + 1):
    rf = fill(BL_LT) if r % 2 == 0 else fill(WHITE)
    for c in range(1, 11):
        cell = ws.cell(row=r, column=c)
        if c in VCOL:
            cell.fill = fill("E2EFDA") if r % 2 == 0 else fill("F0F7EC")
        elif c in PCOL:
            cell.fill = fill("FFF2CC") if r % 2 == 0 else fill("FFFBEE")
        else:
            cell.fill = rf
        cell.border = BRD
        cell.font = fnt(size=10, bold=(c == 8))
        cell.alignment = aln("left", "center") if c <= 2 else aln("center", "center")
        if c >= 3:
            cell.number_format = "#,##0"

for col, w in [("A",11),("B",22),("C",14),("D",14),("E",13),
               ("F",12),("G",14),("H",14),("I",13),("J",13)]:
    ws.column_dimensions[col].width = w
ws.freeze_panes = "C4"

# ── Guardar ───────────────────────────────────────────────────────────────────
wb.save(PATH)
print("Formato aplicado correctamente.")
