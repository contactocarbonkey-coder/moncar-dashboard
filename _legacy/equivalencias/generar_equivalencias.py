import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from collections import defaultdict

# ============================================================
# COLORES - paleta azul (consistente con STOCK_MONCAR)
# ============================================================
DARK_BLUE  = "1F3864"
MED_BLUE   = "2E75B6"
LIGHT_BLUE = "DEEAF1"
YELLOW     = "FFF2CC"
GREEN      = "E2EFDA"
WHITE      = "FFFFFF"

MAX_EQUIV   = 8   # columnas EQUIV en TablaEquiv (max real = 5, 8 por margen)
RESULT_ROWS = MAX_EQUIV

# ============================================================
# NORMALIZACION DE MEDIDAS
# Unifica formatos para que "6.00" (string), 6 (int) y 6.0 (float)
# se traten como el mismo valor al buscar equivalencias.
#   6.00  ->  "6"
#   6     ->  "6"
#   6.3   ->  "6,3"
#   6.30  ->  "6,3"
#   7.90  ->  "7,9"
#   DIAM  ->  "DIAM"
# ============================================================
def norm(val):
    if val is None:
        return ""
    s = str(val).strip()
    if s == "":
        return ""
    if s.upper() == "DIAM":
        return "DIAM"
    try:
        f = float(s.replace(",", "."))
        # :.10g elimina ceros al final: 6.00 -> "6", 6.30 -> "6.3"
        result = f"{f:.10g}"
        # Separador decimal a coma (estandar argentino)
        return result.replace(".", ",")
    except (ValueError, OverflowError):
        return s   # si no es numero, lo devuelvo tal cual

# ============================================================
# HELPERS DE ESTILO
# ============================================================
def style_hdr(cell, dark=True):
    cell.font      = Font(bold=True, color=WHITE, name="Calibri", size=11)
    cell.fill      = PatternFill("solid", fgColor=DARK_BLUE if dark else MED_BLUE)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def thin_box(color=DARK_BLUE):
    s = Side(style="thin", color=color)
    return Border(top=s, bottom=s, left=s, right=s)

def style_data(cell, even=False):
    if even:
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    cell.alignment = Alignment(vertical="center", horizontal="left")

# ============================================================
# 1. LEER Y NORMALIZAR DATOS DEL ARCHIVO FUENTE
# ============================================================
print("Leyendo LP 10 CARBONES 01.04.2026.xlsx ...")
wb_lp = openpyxl.load_workbook("LP 10 CARBONES 01.04.2026.xlsx", data_only=True)
ws_src = wb_lp["lista completa"]

all_rows = []    # lista de dicts con valores YA normalizados
productos = {}   # cod -> primer registro (para los VLOOKUP de detalles)

for row in ws_src.iter_rows(min_row=3, values_only=True):
    if not row[0]:
        continue
    cod      = str(row[0]).strip()
    terminal = str(row[2]).strip() if row[2] is not None else ""
    esp      = norm(row[3])
    anch     = norm(row[4])
    larg     = norm(row[5])
    marca    = str(row[6]).strip() if row[6] else ""
    modelos  = str(row[7]).strip() if row[7] else ""
    precio   = row[8]

    all_rows.append({
        "cod": cod, "terminal": terminal, "esp": esp,
        "anch": anch, "larg": larg, "marca": marca,
        "modelos": modelos, "precio": precio,
    })
    if cod not in productos:
        productos[cod] = {
            "terminal": terminal, "esp": esp, "anch": anch,
            "larg": larg, "marca": marca, "modelos": modelos, "precio": precio,
        }

print(f"  -> {len(all_rows)} filas, {len(productos)} codigos unicos")

# ============================================================
# 2. CALCULAR EQUIVALENCIAS
#    Criterio: mismo TERMINAL + ESP + ANCH + LARG (ya normalizados)
# ============================================================
grupos = defaultdict(set)
for d in all_rows:
    if d["terminal"] and d["esp"] and d["anch"] and d["larg"]:
        clave = (d["terminal"], d["esp"], d["anch"], d["larg"])
        grupos[clave].add(d["cod"])

# Tabla pivot: una fila por codigo, columnas EQUIV_1..EQUIV_MAX
tabla_equiv = {}
cods_con_equiv = set()

for clave, cods in grupos.items():
    cods_sorted = sorted(cods)
    if len(cods_sorted) < 2:
        continue
    for cod in cods_sorted:
        equivs = [c for c in cods_sorted if c != cod]
        tabla_equiv[cod] = equivs
        cods_con_equiv.add(cod)

tabla_equiv_rows = sorted(tabla_equiv.items())
max_real = max((len(v) for v in tabla_equiv.values()), default=0)
print(f"  -> {len(tabla_equiv_rows)} codigos con equivalencias, max {max_real} por codigo")

# ============================================================
# 3. CREAR EL ARCHIVO EXCEL
# ============================================================
wb = openpyxl.Workbook()

# ============================================================
# HOJA 1: Buscador
# ============================================================
ws_bus = wb.active
ws_bus.title = "Buscador"

# Fila 1: Titulo
ws_bus.row_dimensions[1].height = 45
t = ws_bus.cell(row=1, column=1, value="BUSCADOR DE EQUIVALENCIAS  |  MONCAR")
t.font      = Font(bold=True, color=WHITE, name="Calibri", size=18)
t.fill      = PatternFill("solid", fgColor=DARK_BLUE)
t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws_bus.merge_cells("A1:H1")

# Fila 2: Instruccion
ws_bus.row_dimensions[2].height = 20
instr = ws_bus.cell(row=2, column=1,
    value="Escribe el codigo en la celda amarilla. "
          "Las equivalencias aparecen automaticamente.")
instr.font      = Font(italic=True, color="555555", name="Calibri", size=10)
instr.fill      = PatternFill("solid", fgColor="EEF4FB")
instr.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws_bus.merge_cells("A2:H2")

# Fila 3: spacer
ws_bus.row_dimensions[3].height = 10

# Fila 4: Input
ws_bus.row_dimensions[4].height = 40
lbl = ws_bus.cell(row=4, column=1, value="CODIGO A BUSCAR:")
lbl.font      = Font(bold=True, name="Calibri", size=13, color=DARK_BLUE)
lbl.alignment = Alignment(horizontal="right", vertical="center")

inp = ws_bus.cell(row=4, column=2)   # B4: el usuario escribe aqui
inp.font      = Font(bold=True, name="Calibri", size=16, color=DARK_BLUE)
inp.fill      = PatternFill("solid", fgColor=YELLOW)
inp.alignment = Alignment(horizontal="center", vertical="center")
inp.border    = thin_box()

hint = ws_bus.cell(row=4, column=3, value="  Ej: CF 001  |  CF 055  |  CF 751  |  CE 044")
hint.font      = Font(italic=True, color="999999", name="Calibri", size=10)
hint.alignment = Alignment(horizontal="left", vertical="center")
ws_bus.merge_cells("C4:H4")

# Fila 5: spacer
ws_bus.row_dimensions[5].height = 8

# Fila 6: Celda de estado
ws_bus.row_dimensions[6].height = 26
status_formula = (
    '=IF($B$4="","Escribe un codigo en la celda amarilla.",'
    'IF(IFERROR(VLOOKUP($B$4,\'Lista de Precios\'!$A:$A,1,0),"")="","CODIGO NO ENCONTRADO en la lista de precios.",'
    'IF(IFERROR(VLOOKUP($B$4,\'Tabla Equiv\'!$A:$A,1,0),"")="","Este codigo NO tiene equivalencias.","")))'
)
st = ws_bus.cell(row=6, column=1, value=status_formula)
st.font      = Font(bold=True, name="Calibri", size=11, color=DARK_BLUE)
st.fill      = PatternFill("solid", fgColor=LIGHT_BLUE)
st.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws_bus.merge_cells("A6:H6")

# Fila 7: Cabeceras de resultados
ws_bus.row_dimensions[7].height = 30
hdrs = ["COD. EQUIVALENTE", "TERMINAL", "ESP", "ANCH", "LARG",
        "MARCA", "MODELOS Y APLICACIONES", "PRECIO UNIT"]
for i, h in enumerate(hdrs, 1):
    style_hdr(ws_bus.cell(row=7, column=i))
    ws_bus.cell(row=7, column=i).value = h

# Filas 8+: resultados via VLOOKUP
# Lista de Precios cols: A=COD(1) B=TERMINAL(2) C=ESP(3) D=ANCH(4)
#                        E=LARG(5) F=MARCA(6) G=MODELOS(7) H=PRECIO(8)
LP_COLS = {2: "TERMINAL", 3: "ESP", 4: "ANCH", 5: "LARG",
           6: "MARCA",    7: "MODELOS",        8: "PRECIO"}

for i in range(RESULT_ROWS):
    row    = 8 + i
    ci     = i + 2       # col index en Tabla Equiv (2=EQUIV_1, 3=EQUIV_2, ...)
    even   = (i % 2 == 0)
    ws_bus.row_dimensions[row].height = 20

    # Col A: codigo equivalente
    f_cod = f"=IFERROR(VLOOKUP($B$4,'Tabla Equiv'!$A:$I,{ci},0),\"\")"
    c = ws_bus.cell(row=row, column=1, value=f_cod)
    c.font      = Font(bold=True, name="Calibri", size=11, color=DARK_BLUE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    if even:
        c.fill = PatternFill("solid", fgColor=LIGHT_BLUE)

    # Cols B-H: detalles del equivalente (VLOOKUP en Lista de Precios)
    for col_dest, lp_idx in enumerate(LP_COLS.keys(), 2):
        f = (f'=IF($A{row}="","",'
             f'IFERROR(VLOOKUP($A{row},\'Lista de Precios\'!$A:$H,{lp_idx},0),""))')
        c2 = ws_bus.cell(row=row, column=col_dest, value=f)
        style_data(c2, even=even)

# Anchos de columna
ws_bus.column_dimensions["A"].width = 20
ws_bus.column_dimensions["B"].width = 12
ws_bus.column_dimensions["C"].width = 10
ws_bus.column_dimensions["D"].width = 10
ws_bus.column_dimensions["E"].width = 10
ws_bus.column_dimensions["F"].width = 22
ws_bus.column_dimensions["G"].width = 52
ws_bus.column_dimensions["H"].width = 15

ws_bus.freeze_panes = "A8"

# ============================================================
# HOJA 2: Tabla Equiv  (base de los VLOOKUP del Buscador)
# ============================================================
ws_te = wb.create_sheet("Tabla Equiv")

te_hdrs = ["COD"] + [f"EQUIV_{n}" for n in range(1, MAX_EQUIV + 1)]
for i, h in enumerate(te_hdrs, 1):
    style_hdr(ws_te.cell(row=1, column=i))
    ws_te.cell(row=1, column=i).value = h
ws_te.row_dimensions[1].height = 28

for r_idx, (cod, equivs) in enumerate(tabla_equiv_rows, 2):
    ws_te.cell(row=r_idx, column=1, value=cod)
    for j, eq in enumerate(equivs, 2):
        ws_te.cell(row=r_idx, column=j, value=eq)
    even = (r_idx % 2 == 0)
    for c in range(1, len(te_hdrs) + 1):
        cell = ws_te.cell(row=r_idx, column=c)
        if even:
            cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center")

for i in range(1, len(te_hdrs) + 1):
    ws_te.column_dimensions[ws_te.cell(row=1, column=i).column_letter].width = 12

ws_te.freeze_panes = "A2"

# ============================================================
# HOJA 3: Lista de Precios  (valores normalizados + flag equiv)
# ============================================================
ws_lp = wb.create_sheet("Lista de Precios")

lp_hdrs = ["COD", "TERMINAL", "ESP", "ANCH", "LARG",
           "MARCA", "MODELOS Y APLICACIONES", "PRECIO UNIT", "TIENE EQUIV."]
for i, h in enumerate(lp_hdrs, 1):
    style_hdr(ws_lp.cell(row=1, column=i))
    ws_lp.cell(row=1, column=i).value = h
ws_lp.row_dimensions[1].height = 28

for r_idx, d in enumerate(all_rows, 2):
    tiene_equiv = "SI" if d["cod"] in cods_con_equiv else ""
    row_vals = [d["cod"], d["terminal"], d["esp"], d["anch"], d["larg"],
                d["marca"], d["modelos"], d["precio"], tiene_equiv]
    even = (r_idx % 2 == 0)
    for c_idx, val in enumerate(row_vals, 1):
        cell = ws_lp.cell(row=r_idx, column=c_idx, value=val)
        cell.alignment = Alignment(vertical="center", horizontal="left")
        if even:
            cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        if c_idx == 9 and val == "SI":
            cell.fill      = PatternFill("solid", fgColor=GREEN)
            cell.font      = Font(bold=True, color="375623", name="Calibri", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")

col_widths_lp = [12, 11, 10, 10, 10, 22, 55, 14, 13]
for i, w in enumerate(col_widths_lp, 1):
    ws_lp.column_dimensions[ws_lp.cell(row=1, column=i).column_letter].width = w

ws_lp.freeze_panes = "A2"

# ============================================================
# GUARDAR
# ============================================================
output = "EQUIVALENCIAS_MONCAR.xlsx"
wb.save(output)

print()
print(f"Listo: {output}")
print(f"  Buscador:        escribi el codigo en B4")
print(f"  Tabla Equiv:     {len(tabla_equiv_rows)} codigos con equivalencias")
print(f"  Lista de Precios:{len(all_rows)} productos (medidas normalizadas)")
