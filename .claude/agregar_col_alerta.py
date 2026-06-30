"""
Agrega columna K (ALERTA STOCK) en la hoja Movimientos de STOCK_MONCAR.xlsx.
No modifica ninguna otra columna, fórmula ni formato existente.
"""
import openpyxl, shutil
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule

PATH = r'C:\Users\franc\OneDrive\Documentos\KeyCarbon\claude MONCAR RMH\STOCK_MONCAR.xlsx'
BAK  = r'C:\Users\franc\OneDrive\Documentos\KeyCarbon\claude MONCAR RMH\STOCK_MONCAR_preK.xlsx'

shutil.copy2(PATH, BAK)
print("Backup creado.")

# ── Estilos (idénticos a los del resto de la hoja) ────────────────────────────
NAVY  = "1F3864";  BLUE  = "2E75B6"
BL_LT = "DEEAF1";  WHITE = "FFFFFF"

def fill(h):  return PatternFill("solid", fgColor=h)
def fnt(color="1F1F1F", bold=False, size=10):
    return Font(color=color, bold=bold, size=size, name="Calibri")
def aln(h="center", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

_s  = Side(style="thin", color="C0C0C0")
BRD = Border(left=_s, right=_s, top=_s, bottom=_s)

wb = openpyxl.load_workbook(PATH)
ws = wb["Movimientos"]

# ── Fila 1 — extender el fondo del título hasta K ────────────────────────────
ws["K1"].fill = fill(NAVY)
ws["K1"].border = BRD

# ── Fila 2 — extender el fondo de la instrucción hasta K ─────────────────────
ws["K2"].fill = fill(BL_LT)
ws["K2"].border = BRD

# ── Fila 3 — cabecera ─────────────────────────────────────────────────────────
ws["K3"].value     = "ALERTA\nSTOCK"
ws["K3"].fill      = fill(NAVY)
ws["K3"].font      = fnt("FFFFFF", bold=True, size=10)
ws["K3"].alignment = aln("center", "center", wrap=True)
ws["K3"].border    = BRD

# ── Filas 4–1003 — fórmula + estilo ──────────────────────────────────────────
# Lógica:
#   • Vacío  → no es salida o no hay cantidad/COD
#   • "Embolsar: X uds"    → stock embolsado negativo pero hay sin embolsar suficiente
#   • "SIN STOCK: falta X uds" → stock negativo y sin embolsar no alcanza
#
# VLOOKUP a Stock!$A:$J: col3=NO_EMB, col4=EMB_DEP, col5=PREM_DEP, col6=JUANI
# No hay referencia circular: la col K NO es leída por las fórmulas de la hoja Stock.

SALIDAS = 'OR(C{r}="Salida Lita",C{r}="Salida Venta Franco",C{r}="Entrega a Juani",C{r}="Salida Venta Juani")'
VLK_EMB  = 'IFERROR(VLOOKUP(E{r},Stock!$A:$J,IF(D{r}="Premium",5,4),0),0)'
VLK_NOEMB= 'IFERROR(VLOOKUP(E{r},Stock!$A:$J,3,0),0)'
VLK_JUANI= 'IFERROR(VLOOKUP(E{r},Stock!$A:$J,6,0),0)'

def formula(r):
    s  = SALIDAS.replace('{r}', str(r))
    e  = VLK_EMB.replace('{r}', str(r))
    n  = VLK_NOEMB.replace('{r}', str(r))
    j  = VLK_JUANI.replace('{r}', str(r))
    return (
        f'=IF(AND({s},F{r}<>"",E{r}<>""),'
        f'IF(C{r}="Salida Venta Juani",'
            f'IF({j}>=0,"",- {j}&" sin stock Juani"),'
            f'IF({e}>=0,"",'
                f'IF({n}>=-({e}),'
                    f'"Embolsar: "&-({e})&" uds",'
                    f'"SIN STOCK: falta "&-({e})&" uds"'
                f')'
            f')'
        f'),'
        f'"")'
    )

for r in range(4, 1004):
    cell = ws.cell(row=r, column=11)   # columna K
    cell.value     = formula(r)
    cell.border    = BRD
    cell.font      = fnt(size=10)
    cell.alignment = aln("left", "center")
    cell.fill      = fill(BL_LT) if r % 2 == 0 else fill(WHITE)

ws.column_dimensions["K"].width = 26

# ── Formato condicional en K4:K1003 (namespace estándar: Rule+DifferentialStyle)
# Fórmulas simples que solo leen la misma columna K → sin dependencias cruzadas
orange_dxf = DifferentialStyle(fill=PatternFill("solid", fgColor="FAD7B4"))
red_dxf    = DifferentialStyle(fill=PatternFill("solid", fgColor="F2B8B8"))

# priority=1 → mayor prioridad (rojo) se evalúa primero; luego naranja
ws.conditional_formatting.add(
    "K4:K1003",
    Rule(type="expression", dxf=red_dxf,
         formula=['LEFT(K4,3)="SIN"'], priority=1)
)
ws.conditional_formatting.add(
    "K4:K1003",
    Rule(type="expression", dxf=orange_dxf,
         formula=['LEFT(K4,8)="Embolsar"'], priority=2)
)

wb.save(PATH)
print("Columna K agregada correctamente.")
print("Abrí el archivo en Excel y presioná Ctrl+Shift+F9 para recalcular.")
