# -*- coding: utf-8 -*-
"""
generar_sitio.py — genera sitio/index.html con homepage + catalogo MONCAR + RMH.
"""
import os, json
import openpyxl
from collections import Counter

DIR      = os.path.dirname(os.path.abspath(__file__))
OUT_DIR  = os.path.join(DIR, 'sitio')
OUT_PATH = os.path.join(OUT_DIR, 'index.html')

LP_PATH        = os.path.join(DIR, 'LP 10 CARBONES 01.04.2026.xlsx')
STOCK_MC_PATH  = os.path.join(DIR, 'STOCK_MONCAR.xlsx')
STOCK_RMH_PATH = os.path.join(DIR, 'STOCK_RMH.xlsx')

os.makedirs(OUT_DIR, exist_ok=True)

MARCAS_CONFIG = {
    'Bosch':          {'bg': '#007BC0', 'text': '#ffffff'},
    'De Walt':        {'bg': '#FFCD00', 'text': '#000000'},
    'Makita':         {'bg': '#00A8E0', 'text': '#ffffff'},
    'Black Y Decker': {'bg': '#E31837', 'text': '#ffffff'},
    'Metabo':         {'bg': '#003082', 'text': '#ffffff'},
    'Hitachi':        {'bg': '#E4002B', 'text': '#ffffff'},
    'Stanley':        {'bg': '#FFC72C', 'text': '#000000'},
    'Gamma':          {'bg': '#FF6B00', 'text': '#ffffff'},
    'Milwaukee':      {'bg': '#CC0000', 'text': '#ffffff'},
    'Skil':           {'bg': '#003087', 'text': '#ffffff'},
    'Dremel':         {'bg': '#1BA0E1', 'text': '#ffffff'},
    'Einhell':        {'bg': '#E20000', 'text': '#ffffff'},
}

CAT_CONFIG = {
    'Amoladoras':         {'grad': 'linear-gradient(135deg,#FF6B00,#cc3300)', 'icon': '⚙️'},
    'Taladros':           {'grad': 'linear-gradient(135deg,#2563EB,#1e40af)', 'icon': '🔩'},
    'Lijadoras':          {'grad': 'linear-gradient(135deg,#0d9488,#0f766e)', 'icon': '🔧'},
    'Sierras':            {'grad': 'linear-gradient(135deg,#16a34a,#15803d)', 'icon': '🪚'},
    'Lavarropas':         {'grad': 'linear-gradient(135deg,#7c3aed,#6d28d9)', 'icon': '🫧'},
    'Hidrolavadoras':     {'grad': 'linear-gradient(135deg,#0284c7,#0369a1)', 'icon': '💧'},
    'Aspiradoras':        {'grad': 'linear-gradient(135deg,#dc2626,#b91c1c)', 'icon': '🌀'},
    'Electrodomesticos':  {'grad': 'linear-gradient(135deg,#d97706,#b45309)', 'icon': '🔌'},
    'Equipos de Pintura': {'grad': 'linear-gradient(135deg,#db2777,#be185d)', 'icon': '🎨'},
    'Maq. de Coser':      {'grad': 'linear-gradient(135deg,#059669,#047857)', 'icon': '🧵'},
    'Otros':              {'grad': 'linear-gradient(135deg,#374151,#1f2937)', 'icon': '🛠️'},
}

# ---------------------------------------------------------------------------
# Deteccion de categoria
# ---------------------------------------------------------------------------
def detectar_categoria(desc, cod, marca_brand=''):
    d = desc.upper()
    m = marca_brand.upper()

    if cod.startswith('CE'):
        if 'HIDRO' in d:                              return 'Hidrolavadoras'
        if 'LAVARROPAS' in d or 'WHIRLPOOL' in d:    return 'Lavarropas'
        if 'COSER' in d or 'SINGER' in d or 'NECCHI' in d: return 'Maq. de Coser'
        if 'ASPIRADORA' in d or 'TURBION' in d or 'BARRE' in d or 'LUSTRA' in d: return 'Aspiradoras'
        if 'SIERRA' in d or 'LAZO' in d:             return 'Sierras'
        return 'Electrodomesticos'

    if 'HIDROLAVADORA' in m or 'HIDROLAV' in m or 'HIDROLAVADORA' in d: return 'Hidrolavadoras'
    if 'LAVARROPAS' in m or 'LAVARROPAS' in d:       return 'Lavarropas'
    if 'LAVADORA' in d and 'HIDRO' not in d:          return 'Lavarropas'
    if 'ASPIRADORA' in m or 'ASPIRADORAS' in m or 'ASPIRADORA' in d: return 'Aspiradoras'
    if 'BARRE ALFOMBRA' in m or 'BARRE ALFOMBRA' in d: return 'Aspiradoras'
    if 'AMOLADORA' in d or 'AMOL.' in d or 'ESMERIL' in d: return 'Amoladoras'
    if 'AMOLADORA' in m or 'AMOL.' in m:             return 'Amoladoras'
    if 'TALADRO' in d or 'TAL.' in d or 'TAL ' in d: return 'Taladros'
    if 'LIJADORA' in d or 'LIJADORAS' in d or 'LIJA' in d or 'PULIDORA' in d or 'CEPILLO' in d: return 'Lijadoras'
    if 'SIERRA' in d or 'CALADORA' in d or 'TRONZADORA' in d: return 'Sierras'
    if 'PINTURA' in m or 'PINTURA' in d or 'COMPRESOR' in d: return 'Equipos de Pintura'
    return 'Otros'

# ---------------------------------------------------------------------------
# Lectores de Excel
# ---------------------------------------------------------------------------
def leer_stock_moncar():
    """Devuelve dict COD -> {stock: bool, marca: str}"""
    wb = openpyxl.load_workbook(STOCK_MC_PATH, data_only=True)
    ws = wb['Productos']
    result = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        cod = row[0]
        if not cod or not str(cod).startswith(('CC', 'CF', 'CE')):
            continue
        cod   = str(cod).strip()
        marca = str(row[1]).strip().title() if row[1] else 'Moncar'
        vals  = [v for v in row[3:7] if isinstance(v, (int, float))]
        result[cod] = {'stock': sum(vals) > 0, 'marca': marca}
    return result

def leer_rmh():
    wb = openpyxl.load_workbook(STOCK_RMH_PATH, data_only=True)
    ws = wb['Productos']
    productos = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        cod = row[0]
        if not cod or not str(cod).strip().startswith('RM'):
            continue
        cod    = str(cod).strip()
        desc   = str(row[1]).strip() if row[1] else ''
        precio = row[2] if isinstance(row[2], (int, float)) else 0
        vals   = [v for v in row[3:7] if isinstance(v, (int, float))]
        cat    = detectar_categoria(desc, cod)
        productos.append({
            'cod':      cod,
            'linea':    'RMH',
            'marca':    'RMH',
            'cat':      cat,
            'desc':     desc,
            'dims':     '',
            'esp':      None,
            'anch':     None,
            'larg':     None,
            'terminal': None,
            'precio':   int(round(precio)),
            'stock':    sum(vals) > 0,
        })
    return productos

def _es_separador(row):
    """Fila separadora de marca: col0=None, col1=None, col5=nombre (todas mayusc corto)"""
    return (row[0] is None
            and row[1] is None
            and isinstance(row[5], str)
            and 0 < len(row[5].strip()) <= 50
            and row[5].strip().upper() == row[5].strip()
            and row[2] is None)

def _val(v):
    """Convierte dimension a float o None"""
    try:
        return float(str(v).strip())
    except:
        return None

def leer_moncar_lp(stock_map):
    wb = openpyxl.load_workbook(LP_PATH, data_only=True)
    productos = []

    # --- Hoja Precios de Carbones ---
    ws = wb['Precios de Carbones']
    marca_lp = 'Moncar'
    for row in ws.iter_rows(min_row=8, values_only=True):
        if all(v is None for v in row):
            continue
        if _es_separador(row):
            marca_lp = row[5].strip().title()
            continue
        cod = row[0]
        if not cod or not str(cod).strip().startswith(('CF', 'CC')):
            continue
        cod    = str(cod).strip()
        precio = row[6] if isinstance(row[6], (int, float)) and row[6] > 0 else 0
        desc   = str(row[5]).strip() if row[5] else ''
        esp    = _val(row[2])
        anch   = _val(row[3])
        larg   = _val(row[4])
        term   = int(row[1]) if isinstance(row[1], (int, float)) else None
        dims   = 'x'.join(str(v) for v in [row[2], row[3], row[4]]
                          if v is not None and str(v).strip() not in ('', 'None', 'DIAM'))

        # Marca: preferir STOCK_MONCAR (tiene info correcta), fallback a LP
        info   = stock_map.get(cod, {})
        marca  = info.get('marca', marca_lp) if info else marca_lp
        cat    = detectar_categoria(desc, cod, marca)

        productos.append({
            'cod':      cod,
            'linea':    'MONCAR',
            'marca':    marca,
            'cat':      cat,
            'desc':     desc,
            'dims':     dims,
            'esp':      esp,
            'anch':     anch,
            'larg':     larg,
            'terminal': term,
            'precio':   int(round(precio)),
            'stock':    info.get('stock', False),
        })

    # --- Hoja Carbones Electrodomesticos ---
    ws2 = wb['Carbones Electrodomesticos']
    for row in ws2.iter_rows(min_row=2, values_only=True):
        cod = row[0]
        if not cod or not str(cod).strip().startswith('CE'):
            continue
        cod    = str(cod).strip()
        precio = row[7] if isinstance(row[7], (int, float)) and row[7] > 0 else 0
        desc   = str(row[6]).strip() if row[6] else ''
        esp    = _val(row[3])
        anch   = _val(row[4])
        larg   = _val(row[5])
        dims   = 'x'.join(str(v) for v in [row[3], row[4], row[5]]
                          if v is not None and str(v).strip() not in ('', 'None'))
        info   = stock_map.get(cod, {})
        cat    = detectar_categoria(desc, cod)
        productos.append({
            'cod':      cod,
            'linea':    'MONCAR',
            'marca':    'Electrodomesticos',
            'cat':      cat,
            'desc':     desc,
            'dims':     dims,
            'esp':      esp,
            'anch':     anch,
            'larg':     larg,
            'terminal': None,
            'precio':   int(round(precio)),
            'stock':    info.get('stock', False),
        })

    return productos

# ---------------------------------------------------------------------------
# Generador HTML
# ---------------------------------------------------------------------------
def generar_html(productos):
    cat_count = Counter(p['cat'] for p in productos)
    cats_orden = [c for c, _ in sorted(cat_count.items(), key=lambda x: -x[1])]

    EXCLUIR_MARCAS = {'Moncar','Electrodomesticos','Cobreados','Porta Carbones',
                      'Power Tools','Experto','Bta','Neo','Lazo Asincron','Versa',
                      'Lavarropas','Aspiradora','Ventilador','Batidora',
                      'Maquina De Coser','Lustra Aspiradora','Hidrolavadora',
                      'Enceradoras','Afeitadoras','Bordeadora','Barre Alfombras',
                      'Salkor','Stingray','Dover','Gladiator','Stayer','Peugeot',
                      'Dowen Paggio','Martins','Agro','Argentec','Adiabatic'}

    marca_count = Counter(
        p['marca'] for p in productos
        if p['linea'] == 'MONCAR' and p['marca'] not in EXCLUIR_MARCAS
    )
    marcas_destacadas = [(m, cnt) for m, cnt in marca_count.items() if m in MARCAS_CONFIG]
    marcas_destacadas.sort(key=lambda x: -x[1])

    # Datos para filtros del sidebar
    todas_marcas   = sorted(m for m, _ in marcas_destacadas)
    todas_cats     = sorted(cat_count.keys())
    precios_validos = [p['precio'] for p in productos if p['precio'] > 0]
    precio_min     = min(precios_validos) if precios_validos else 0
    precio_max     = max(precios_validos) if precios_validos else 99999

    esps  = sorted(set(p['esp']  for p in productos if p['esp']  is not None))
    anchs = sorted(set(p['anch'] for p in productos if p['anch'] is not None))
    largs = sorted(set(p['larg'] for p in productos if p['larg'] is not None))
    terms = sorted(set(p['terminal'] for p in productos if p['terminal'] is not None))

    def cat_cards():
        cards = []
        for cat in cats_orden:
            cfg   = CAT_CONFIG.get(cat, CAT_CONFIG['Otros'])
            count = cat_count[cat]
            cards.append(
                f'<div class="cat-card" style="background:{cfg["grad"]}" '
                f'onclick="irACat(\'{cat}\')">'
                f'<div class="cat-icon">{cfg["icon"]}</div>'
                f'<div class="cat-name">{cat}</div>'
                f'<div class="cat-count">{count} productos</div>'
                f'</div>'
            )
        return '\n'.join(cards)

    def marca_cards():
        cards = []
        for m, cnt in marcas_destacadas:
            cfg = MARCAS_CONFIG[m]
            cards.append(
                f'<div class="marca-card" style="background:{cfg["bg"]};color:{cfg["text"]}" '
                f'onclick="irAMarca(\'{m}\')">'
                f'<div class="marca-name">{m}</div>'
                f'<div class="marca-sub" style="color:{cfg["text"]};opacity:.72">'
                f'{cnt} carbones</div>'
                f'</div>'
            )
        return '\n'.join(cards)

    def sidebar_marcas():
        opts = []
        for m, cnt in sorted(marca_count.items(), key=lambda x: -x[1]):
            opts.append(
                f'<label class="sb-check">'
                f'<input type="checkbox" value="{m}" onchange="filtrar()">'
                f'<span>{m}</span><span class="sb-cnt">{cnt}</span>'
                f'</label>'
            )
        return '\n'.join(opts)

    def sidebar_cats():
        opts = []
        for cat in cats_orden:
            opts.append(
                f'<label class="sb-check">'
                f'<input type="checkbox" value="{cat}" onchange="filtrar()">'
                f'<span>{cat}</span><span class="sb-cnt">{cat_count[cat]}</span>'
                f'</label>'
        )
        return '\n'.join(opts)

    def opt(vals):
        return '\n'.join(f'<option value="{v}">{v}</option>' for v in vals)

    productos_json = json.dumps(productos, ensure_ascii=False)

    return f'''<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>MONCAR RMH — Repuestos para Herramientas</title>
<style>
*{{box-sizing:border-box;margin:0;padding:0}}
:root{{
  --bg:#f4f4f4;--white:#fff;--surface:#fff;--border:#e2e2e2;
  --orange:#FF6B00;--orange2:#e05e00;
  --text:#111;--muted:#666;
  --green:#16a34a;--red:#dc2626;
  --nav-h:64px;--sb-w:260px;
}}
html{{scroll-behavior:smooth}}
body{{background:var(--bg);color:var(--text);font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,sans-serif}}

/* NAV */
nav{{position:fixed;top:0;left:0;right:0;z-index:100;background:var(--white);border-bottom:1px solid var(--border);height:var(--nav-h);display:flex;align-items:center;padding:0 28px;box-shadow:0 1px 4px rgba(0,0,0,.06)}}
.nav-logo{{font-size:20px;font-weight:900;letter-spacing:.5px;cursor:pointer;margin-right:28px}}
.nav-logo span{{color:var(--orange)}}
.nav-links{{display:flex;gap:2px}}
.nav-link{{color:var(--muted);font-size:14px;font-weight:500;cursor:pointer;padding:6px 14px;border-radius:6px;border:none;background:none;transition:.15s}}
.nav-link:hover{{color:var(--text);background:#f0f0f0}}
.nav-link.active{{color:var(--orange);font-weight:700}}
.nav-right{{display:flex;gap:10px;align-items:center;margin-left:auto}}
.btn-wa-nav{{background:#25D366;color:#fff;border:none;padding:8px 16px;border-radius:7px;font-size:13px;font-weight:700;cursor:pointer;text-decoration:none;display:flex;align-items:center;gap:5px}}
.btn-cart-nav{{background:var(--orange);color:#fff;border:none;padding:8px 16px;border-radius:7px;font-size:13px;font-weight:700;cursor:pointer;display:flex;align-items:center;gap:6px}}
.btn-cart-nav:hover{{background:var(--orange2)}}
#cart-count{{background:#fff;color:var(--orange);border-radius:50%;width:18px;height:18px;display:inline-flex;align-items:center;justify-content:center;font-size:11px;font-weight:900}}

/* PAGES */
.page{{display:none;padding-top:var(--nav-h)}}
.page.active{{display:block}}

/* ===== INICIO ===== */
.hero{{background:linear-gradient(135deg,#111 0%,#1a0800 60%,#111 100%);padding:72px 28px 64px;text-align:center;position:relative;overflow:hidden}}
.hero::before{{content:'';position:absolute;inset:0;background:radial-gradient(ellipse at 55% 50%,rgba(255,107,0,.18) 0%,transparent 65%)}}
.hero-inner{{position:relative;z-index:1;max-width:640px;margin:0 auto}}
.hero-badge{{display:inline-block;background:rgba(255,107,0,.15);color:#FF6B00;border:1px solid rgba(255,107,0,.35);border-radius:20px;padding:5px 16px;font-size:12px;font-weight:700;letter-spacing:1px;text-transform:uppercase;margin-bottom:18px}}
.hero h1{{font-size:clamp(30px,5vw,50px);font-weight:900;line-height:1.1;color:#fff;margin-bottom:14px}}
.hero h1 em{{color:var(--orange);font-style:normal}}
.hero p{{color:rgba(255,255,255,.65);font-size:16px;margin-bottom:28px;line-height:1.6}}
.hero-btns{{display:flex;gap:12px;justify-content:center;flex-wrap:wrap}}
.btn-hero-p{{background:var(--orange);color:#fff;border:none;padding:13px 28px;border-radius:8px;font-size:15px;font-weight:700;cursor:pointer}}
.btn-hero-p:hover{{background:var(--orange2)}}
.btn-hero-s{{color:#fff;border:1px solid rgba(255,255,255,.25);padding:13px 28px;border-radius:8px;font-size:15px;font-weight:600;cursor:pointer;text-decoration:none;display:inline-block;background:transparent}}
.btn-hero-s:hover{{border-color:rgba(255,255,255,.5)}}

.sec{{padding:52px 28px}}
.sec-alt{{background:var(--white)}}
.sec-head{{margin-bottom:28px}}
.sec-head h2{{font-size:clamp(22px,3vw,28px);font-weight:900}}
.sec-head p{{color:var(--muted);font-size:14px;margin-top:5px}}

/* MARCA CARDS */
.marca-grid{{display:grid;grid-template-columns:repeat(auto-fill,minmax(155px,1fr));gap:14px}}
.marca-card{{border-radius:12px;padding:22px 16px;cursor:pointer;transition:transform .15s,box-shadow .15s;text-align:center}}
.marca-card:hover{{transform:translateY(-3px);box-shadow:0 8px 20px rgba(0,0,0,.2)}}
.marca-name{{font-size:16px;font-weight:800;letter-spacing:.3px}}
.marca-sub{{font-size:12px;margin-top:4px}}

/* CAT CARDS */
.cat-grid{{display:grid;grid-template-columns:repeat(auto-fill,minmax(155px,1fr));gap:14px}}
.cat-card{{border-radius:12px;padding:22px 16px;cursor:pointer;transition:transform .15s,box-shadow .15s;text-align:center}}
.cat-card:hover{{transform:translateY(-3px);box-shadow:0 8px 20px rgba(0,0,0,.15)}}
.cat-icon{{font-size:32px;margin-bottom:10px}}
.cat-name{{color:#fff;font-size:14px;font-weight:700;text-transform:uppercase;letter-spacing:.5px}}
.cat-count{{color:rgba(255,255,255,.75);font-size:12px;margin-top:3px}}

/* NOSOTROS */
.nosotros-wrap{{max-width:900px;margin:0 auto}}
.nos-hero{{background:linear-gradient(135deg,#111,#1a0800);border-radius:16px;padding:48px 40px;margin-bottom:36px;text-align:center}}
.nos-hero h2{{font-size:30px;font-weight:900;color:#fff;margin-bottom:12px}}
.nos-hero h2 span{{color:var(--orange)}}
.nos-hero p{{color:rgba(255,255,255,.65);font-size:16px;line-height:1.7;max-width:560px;margin:0 auto}}
.nos-cards{{display:grid;grid-template-columns:repeat(auto-fit,minmax(200px,1fr));gap:16px;margin-bottom:28px}}
.n-card{{background:var(--white);border:1px solid var(--border);border-radius:12px;padding:22px}}
.n-icon{{font-size:26px;margin-bottom:10px}}
.n-card h3{{font-size:14px;font-weight:800;margin-bottom:6px}}
.n-card p{{font-size:13px;color:var(--muted);line-height:1.6}}
.contact-box{{background:var(--white);border:1px solid var(--border);border-radius:12px;padding:24px}}
.contact-box h3{{font-size:15px;font-weight:800;margin-bottom:16px}}
.c-row{{display:flex;align-items:center;gap:14px;padding:11px 0;border-bottom:1px solid var(--border)}}
.c-row:last-child{{border-bottom:none}}
.c-ico{{width:38px;height:38px;border-radius:9px;display:flex;align-items:center;justify-content:center;font-size:17px;flex-shrink:0}}
.c-lbl{{font-size:11px;color:var(--muted);font-weight:700;text-transform:uppercase;letter-spacing:.5px}}
.c-val{{font-size:14px;font-weight:700}}
.c-val a{{color:var(--text);text-decoration:none}}
.c-val a:hover{{color:var(--orange)}}

/* ===== PRODUCTOS - layout con sidebar ===== */
.prod-layout{{display:flex;align-items:flex-start;min-height:calc(100vh - var(--nav-h))}}

/* SIDEBAR */
.sidebar{{width:var(--sb-w);flex-shrink:0;background:var(--white);border-right:1px solid var(--border);min-height:calc(100vh - var(--nav-h));position:sticky;top:var(--nav-h);max-height:calc(100vh - var(--nav-h));overflow-y:auto;padding:20px 16px}}
.sb-title{{font-size:13px;font-weight:800;text-transform:uppercase;letter-spacing:.6px;color:var(--muted);margin-bottom:12px;margin-top:4px}}
.sb-section{{margin-bottom:20px;border-bottom:1px solid var(--border);padding-bottom:18px}}
.sb-section:last-child{{border-bottom:none}}
.sb-check{{display:flex;align-items:center;gap:7px;font-size:13px;cursor:pointer;padding:3px 0;color:var(--text)}}
.sb-check input{{accent-color:var(--orange);width:14px;height:14px;cursor:pointer}}
.sb-check span:first-of-type{{flex:1}}
.sb-cnt{{color:var(--muted);font-size:11px}}
.sb-linea{{display:flex;gap:6px;flex-wrap:wrap}}
.sb-linea-btn{{border:1px solid var(--border);background:var(--white);color:var(--muted);border-radius:6px;padding:5px 12px;font-size:12px;font-weight:600;cursor:pointer;transition:.15s}}
.sb-linea-btn:hover{{background:var(--bg)}}
.sb-linea-btn.active{{background:var(--orange);border-color:var(--orange);color:#fff}}
.sb-range{{display:flex;gap:8px;align-items:center}}
.sb-range input{{width:80px;background:var(--bg);border:1px solid var(--border);border-radius:6px;padding:5px 8px;font-size:13px;color:var(--text)}}
.sb-range span{{color:var(--muted);font-size:13px}}
.sb-select{{width:100%;background:var(--bg);border:1px solid var(--border);border-radius:6px;padding:6px 8px;font-size:13px;color:var(--text);margin-bottom:6px}}
.sb-stock{{display:flex;align-items:center;gap:8px;font-size:13px;cursor:pointer}}
.sb-stock input{{accent-color:var(--orange);width:15px;height:15px}}
.sb-clear{{width:100%;background:none;border:1px solid var(--border);border-radius:7px;padding:8px;font-size:13px;color:var(--muted);cursor:pointer;margin-top:4px}}
.sb-clear:hover{{border-color:var(--orange);color:var(--orange)}}

/* MAIN CONTENT */
.prod-main{{flex:1;min-width:0}}
.prod-topbar{{padding:16px 20px;background:var(--white);border-bottom:1px solid var(--border);display:flex;align-items:center;gap:12px}}
.search-wrap{{position:relative;flex:1}}
.search-wrap input{{width:100%;background:var(--bg);border:1px solid var(--border);border-radius:8px;padding:9px 36px 9px 12px;color:var(--text);font-size:14px}}
.search-wrap input:focus{{outline:none;border-color:var(--orange)}}
.search-wrap input::placeholder{{color:var(--muted)}}
.search-clear{{position:absolute;right:10px;top:50%;transform:translateY(-50%);background:none;border:none;color:var(--muted);cursor:pointer;font-size:15px;display:none}}
.active-filters{{padding:8px 20px;background:var(--bg);display:none;gap:8px;flex-wrap:wrap}}
.active-tag{{background:rgba(255,107,0,.1);border:1px solid rgba(255,107,0,.25);color:var(--orange);border-radius:20px;padding:3px 12px;font-size:12px;font-weight:600;display:flex;align-items:center;gap:5px;cursor:pointer}}
.grid-wrap{{padding:20px}}
.results-info{{font-size:13px;color:var(--muted);margin-bottom:14px}}
.grid{{display:grid;grid-template-columns:repeat(auto-fill,minmax(195px,1fr));gap:13px}}
.no-results{{text-align:center;padding:60px 20px;color:var(--muted)}}

/* PRODUCT CARD */
.card{{background:var(--white);border:1px solid var(--border);border-radius:10px;overflow:hidden;display:flex;flex-direction:column;transition:border-color .15s,box-shadow .15s}}
.card:hover{{border-color:var(--orange);box-shadow:0 4px 16px rgba(255,107,0,.1)}}
.card-icon{{background:var(--bg);display:flex;align-items:center;justify-content:center;padding:18px;height:88px}}
.card-body{{padding:11px;flex:1;display:flex;flex-direction:column;gap:4px}}
.card-badges{{display:flex;gap:4px;flex-wrap:wrap}}
.badge{{display:inline-block;font-size:10px;border-radius:4px;padding:2px 6px;font-weight:700}}
.bl{{background:rgba(255,107,0,.1);color:var(--orange)}}
.bc{{background:#f0f0f0;color:var(--muted)}}
.card-cod{{font-size:11px;font-weight:700;color:var(--orange);letter-spacing:.5px;margin-top:2px}}
.card-desc{{font-size:13px;line-height:1.35;flex:1}}
.card-sub{{font-size:11px;color:var(--muted)}}
.card-precio{{font-size:16px;font-weight:900;margin-top:4px}}
.card-stock{{font-size:11px;font-weight:600;padding:2px 7px;border-radius:4px;align-self:flex-start}}
.sok{{background:rgba(22,163,74,.1);color:var(--green)}}
.sno{{background:rgba(220,38,38,.1);color:var(--red)}}
.btn-add{{margin-top:7px;background:var(--orange);color:#fff;border:none;border-radius:6px;padding:8px;font-size:13px;font-weight:600;cursor:pointer;width:100%}}
.btn-add:hover{{background:var(--orange2)}}
.btn-add:disabled{{background:#e5e5e5;color:#aaa;cursor:not-allowed}}

/* CART */
.cart-overlay{{display:none;position:fixed;inset:0;background:rgba(0,0,0,.4);z-index:200}}
.cart-panel{{position:fixed;top:0;right:-420px;width:min(420px,100vw);height:100%;background:var(--white);border-left:1px solid var(--border);z-index:201;display:flex;flex-direction:column;transition:right .25s;box-shadow:-4px 0 20px rgba(0,0,0,.1)}}
.cart-panel.open{{right:0}}
.cart-header{{padding:18px 20px;border-bottom:1px solid var(--border);display:flex;align-items:center;justify-content:space-between}}
.cart-header h2{{font-size:16px;font-weight:800}}
.btn-xc{{background:none;border:none;color:var(--muted);font-size:22px;cursor:pointer;line-height:1}}
.cart-items{{flex:1;overflow-y:auto;padding:12px}}
.cart-empty{{text-align:center;padding:40px 20px;color:var(--muted)}}
.ci{{display:flex;gap:10px;align-items:flex-start;padding:10px 0;border-bottom:1px solid var(--border)}}
.ci-info{{flex:1}}
.ci-cod{{font-size:11px;color:var(--orange);font-weight:700}}
.ci-desc{{font-size:13px;margin:2px 0}}
.ci-precio{{font-size:14px;font-weight:700}}
.ci-qty{{display:flex;align-items:center;gap:6px;margin-top:5px}}
.ci-qty button{{background:#f0f0f0;border:none;width:24px;height:24px;border-radius:4px;cursor:pointer;font-size:15px}}
.ci-qty span{{font-size:14px;font-weight:700;min-width:20px;text-align:center}}
.btn-rm{{background:none;border:none;color:var(--red);cursor:pointer;font-size:17px;padding:0 4px;line-height:1}}
.cart-footer{{padding:16px 20px;border-top:1px solid var(--border)}}
.cart-total{{display:flex;justify-content:space-between;font-size:16px;font-weight:800;margin-bottom:14px}}
.btn-wa-c{{background:#25D366;color:#fff;border:none;border-radius:8px;padding:12px;font-size:15px;font-weight:700;cursor:pointer;width:100%;margin-bottom:8px}}
.btn-em-c{{background:var(--white);color:var(--text);border:1px solid var(--border);border-radius:8px;padding:12px;font-size:14px;font-weight:600;cursor:pointer;width:100%}}
.btn-em-c:hover{{border-color:var(--orange);color:var(--orange)}}

/* FOOTER */
footer{{background:var(--white);border-top:1px solid var(--border);padding:36px 28px 24px}}
.fi{{display:grid;grid-template-columns:repeat(auto-fit,minmax(180px,1fr));gap:24px;max-width:1000px;margin:0 auto}}
.fb{{font-size:18px;font-weight:900;margin-bottom:6px}}
.fb span{{color:var(--orange)}}
.fd{{font-size:13px;color:var(--muted);line-height:1.5}}
.fc h4{{font-size:11px;font-weight:700;text-transform:uppercase;letter-spacing:.6px;color:var(--muted);margin-bottom:10px}}
.fc a{{display:block;font-size:13px;color:var(--text);text-decoration:none;margin-bottom:5px;cursor:pointer}}
.fc a:hover{{color:var(--orange)}}
.ft{{text-align:center;padding-top:18px;margin-top:20px;border-top:1px solid var(--border);font-size:12px;color:var(--muted)}}

/* MOBILE SIDEBAR TOGGLE */
.sb-toggle{{display:none;background:var(--orange);color:#fff;border:none;padding:8px 16px;border-radius:7px;font-size:13px;font-weight:700;cursor:pointer}}
.sidebar-overlay{{display:none;position:fixed;inset:0;background:rgba(0,0,0,.4);z-index:50}}

@media(max-width:860px){{
  .sidebar{{position:fixed;top:var(--nav-h);left:-280px;z-index:51;transition:left .25s;max-height:100vh}}
  .sidebar.open{{left:0}}
  .sidebar-overlay.open{{display:block}}
  .sb-toggle{{display:flex;align-items:center;gap:6px}}
  .prod-main{{width:100%}}
}}
@media(max-width:640px){{
  .sec{{padding:36px 16px}}
  .marca-grid,.cat-grid{{grid-template-columns:repeat(2,1fr)}}
  .nos-hero{{padding:32px 20px}}
  .grid{{grid-template-columns:repeat(2,1fr);gap:10px}}
  .grid-wrap{{padding:14px}}
}}

/* FORM PEDIDO MODAL */
.fmodal-ov{{display:none;position:fixed;inset:0;background:rgba(0,0,0,.5);z-index:400;align-items:center;justify-content:center;padding:20px}}
.fmodal-ov.open{{display:flex}}
.fmodal{{background:var(--white);border-radius:14px;padding:28px;width:min(480px,100%);max-height:90vh;overflow-y:auto;box-shadow:0 20px 60px rgba(0,0,0,.25)}}
.fmodal-title{{font-size:19px;font-weight:900;margin-bottom:5px}}
.fmodal-sub{{font-size:13px;color:var(--muted);margin-bottom:22px;line-height:1.5}}
.fg{{margin-bottom:14px}}
.fg label{{display:block;font-size:13px;font-weight:700;margin-bottom:5px}}
.fg input{{width:100%;border:1px solid var(--border);border-radius:7px;padding:10px 12px;font-size:14px;color:var(--text);background:var(--bg)}}
.fg input:focus{{outline:none;border-color:var(--orange)}}
.fg input.ferr{{border-color:var(--red)}}
.fg-note{{font-size:12px;color:var(--muted);margin-top:3px}}
.zona-res{{border-radius:8px;padding:11px 14px;margin-top:4px;font-size:13px;line-height:1.4;display:none}}
.zok{{background:#f0fdf4;border:1px solid #86efac;color:#166534}}
.zwarn{{background:#fffbeb;border:1px solid #fcd34d;color:#92400e}}
.fmodal-acts{{display:flex;gap:10px;margin-top:22px}}
.btn-cfm{{flex:1;background:var(--orange);color:#fff;border:none;border-radius:8px;padding:12px;font-size:15px;font-weight:700;cursor:pointer}}
.btn-cfm:hover{{background:var(--orange2)}}
.btn-cnc{{background:#f0f0f0;color:var(--muted);border:none;border-radius:8px;padding:12px 16px;font-size:14px;font-weight:600;cursor:pointer}}
</style>
</head>
<body>

<nav>
  <div class="nav-logo" onclick="irAInicio()">MONCAR<span> RMH</span></div>
  <div class="nav-links">
    <button class="nav-link active" id="nav-inicio" onclick="irAInicio()">Inicio</button>
    <button class="nav-link" id="nav-productos" onclick="irAProductos()">Productos</button>
    <button class="nav-link" id="nav-nosotros" onclick="irANosotros()">Nosotros</button>
  </div>
  <div class="nav-right">
    <a class="btn-wa-nav" href="https://wa.me/5491136267009" target="_blank">&#128172; WhatsApp</a>
    <button class="btn-cart-nav" onclick="abrirCart()">Carrito <span id="cart-count">0</span></button>
  </div>
</nav>

<!-- ===== INICIO ===== -->
<div id="page-inicio" class="page active">
  <div class="hero">
    <div class="hero-inner">
      <div class="hero-badge">Repuestos originales</div>
      <h1>Carbones y repuestos<br>para tus <em>herramientas</em></h1>
      <p>Carbones electricos MONCAR e interruptores RMH.<br>Compatible con las principales marcas. Envios a todo el pais.</p>
      <div class="hero-btns">
        <button class="btn-hero-p" onclick="irAProductos()">Ver catalogo completo</button>
        <a class="btn-hero-s" href="https://wa.me/5491136267009" target="_blank">Consultar por WhatsApp</a>
      </div>
    </div>
  </div>

  <div class="sec">
    <div class="sec-head"><h2>Compra por marca</h2><p>Compatible con las principales marcas del mercado</p></div>
    <div class="marca-grid">{marca_cards()}</div>
  </div>

  <div class="sec sec-alt">
    <div class="sec-head"><h2>Compra por categoria</h2><p>Encontra rapido el repuesto para tu herramienta</p></div>
    <div class="cat-grid">{cat_cards()}</div>
  </div>
</div>

<!-- ===== PRODUCTOS ===== -->
<div id="page-productos" class="page">
  <div class="prod-layout">

    <!-- SIDEBAR -->
    <div class="sidebar" id="sidebar">
      <div class="sb-section">
        <div class="sb-title">Linea</div>
        <div class="sb-linea">
          <button class="sb-linea-btn active" data-val="" onclick="setLinea(this)">Todos</button>
          <button class="sb-linea-btn" data-val="MONCAR" onclick="setLinea(this)">MONCAR</button>
          <button class="sb-linea-btn" data-val="RMH" onclick="setLinea(this)">RMH</button>
        </div>
      </div>

      <div class="sb-section">
        <div class="sb-title">Disponibilidad</div>
        <label class="sb-stock">
          <input type="checkbox" id="solo-stock" onchange="filtrar()">
          <span>Solo en stock</span>
        </label>
      </div>

      <div class="sb-section">
        <div class="sb-title">Categoria</div>
        <div id="sb-cats">{sidebar_cats()}</div>
      </div>

      <div class="sb-section">
        <div class="sb-title">Marca</div>
        <div id="sb-marcas">{sidebar_marcas()}</div>
      </div>

      <div class="sb-section">
        <div class="sb-title">Precio</div>
        <div class="sb-range">
          <input type="number" id="precio-min" placeholder="Min" value="" oninput="filtrar()">
          <span>—</span>
          <input type="number" id="precio-max" placeholder="Max" value="" oninput="filtrar()">
        </div>
      </div>

      <div class="sb-section">
        <div class="sb-title">Dimensiones (mm)</div>
        <select class="sb-select" id="fil-esp" onchange="filtrar()">
          <option value="">Espesor: todos</option>
          {opt(esps)}
        </select>
        <select class="sb-select" id="fil-anch" onchange="filtrar()">
          <option value="">Ancho: todos</option>
          {opt(anchs)}
        </select>
        <select class="sb-select" id="fil-larg" onchange="filtrar()">
          <option value="">Largo: todos</option>
          {opt(largs)}
        </select>
      </div>

      <div class="sb-section">
        <div class="sb-title">Terminal</div>
        <select class="sb-select" id="fil-term" onchange="filtrar()">
          <option value="">Todos los terminales</option>
          {opt(terms)}
        </select>
      </div>

      <button class="sb-clear" onclick="limpiarFiltros()">Limpiar filtros</button>
    </div><!-- /sidebar -->

    <div class="sidebar-overlay" id="sb-overlay" onclick="cerrarSidebar()"></div>

    <!-- MAIN -->
    <div class="prod-main">
      <div class="prod-topbar">
        <button class="sb-toggle" onclick="abrirSidebar()">&#9776; Filtros</button>
        <div class="search-wrap">
          <input type="text" id="search" placeholder="Buscar codigo, descripcion, marca..." oninput="filtrar()">
          <button class="search-clear" id="clear-btn" onclick="limpiarBusqueda()">&#x2715;</button>
        </div>
      </div>
      <div class="active-filters" id="active-filters"></div>
      <div class="grid-wrap">
        <div class="results-info" id="results-info"></div>
        <div class="grid" id="grid"></div>
        <div class="no-results" id="no-results" style="display:none">
          <p style="font-size:18px;margin-bottom:8px">Sin resultados</p>
          <span>Proba ajustando los filtros</span>
        </div>
      </div>
    </div>

  </div>
</div>

<!-- ===== NOSOTROS ===== -->
<div id="page-nosotros" class="page">
  <div class="sec">
    <div class="nosotros-wrap">
      <div class="nos-hero">
        <h2>Sobre <span>MONCAR RMH</span></h2>
        <p>Empresa familiar especializada en repuestos electricos para herramientas. Carbones electricos MONCAR e interruptores RMH compatibles con las principales marcas del mercado.</p>
      </div>
      <div class="nos-cards">
        <div class="n-card"><div class="n-icon">⚡</div><h3>Experiencia en el rubro</h3><p>Anos de trayectoria en la venta de repuestos electricos para herramientas profesionales y de uso domestico.</p></div>
        <div class="n-card"><div class="n-icon">🏷️</div><h3>Multimarca</h3><p>Compatible con Bosch, DeWalt, Makita, Black & Decker, Metabo, Hitachi, Stanley, Gamma, Milwaukee y muchas mas.</p></div>
        <div class="n-card"><div class="n-icon">📦</div><h3>Envios a todo el pais</h3><p>Despachamos a cualquier punto del pais. Consulta por tu zona y los tiempos de entrega.</p></div>
        <div class="n-card"><div class="n-icon">💬</div><h3>Atencion personalizada</h3><p>Te ayudamos a encontrar el repuesto exacto para tu herramienta. Consulta sin compromiso por WhatsApp.</p></div>
      </div>
      <div class="contact-box">
        <h3>Contacto</h3>
        <div class="c-row">
          <div class="c-ico" style="background:#dcfce7">&#128172;</div>
          <div><div class="c-lbl">WhatsApp</div><div class="c-val"><a href="https://wa.me/5491136267009" target="_blank">+54 9 11 3626-7009</a></div></div>
        </div>
        <div class="c-row">
          <div class="c-ico" style="background:#fff7ed">&#128140;</div>
          <div><div class="c-lbl">Email</div><div class="c-val"><a href="mailto:contacto.carbonkey@gmail.com">contacto.carbonkey@gmail.com</a></div></div>
        </div>
        <div class="c-row">
          <div class="c-ico" style="background:#eff6ff">&#128200;</div>
          <div><div class="c-lbl">Catalogo</div><div class="c-val"><a onclick="irAProductos()" style="cursor:pointer;color:var(--orange)">{len(productos)} productos disponibles</a></div></div>
        </div>
      </div>
    </div>
  </div>
</div>

<footer>
  <div class="fi">
    <div><div class="fb">MONCAR<span> RMH</span></div><div class="fd">Carbones electricos e interruptores para herramientas. Envios a todo el pais.</div></div>
    <div class="fc"><h4>Navegacion</h4><a onclick="irAInicio()">Inicio</a><a onclick="irAProductos()">Productos</a><a onclick="irANosotros()">Nosotros</a></div>
    <div class="fc"><h4>Contacto</h4><a href="https://wa.me/5491136267009" target="_blank">WhatsApp: +54 9 11 3626-7009</a><a href="mailto:contacto.carbonkey@gmail.com">contacto.carbonkey@gmail.com</a></div>
  </div>
  <div class="ft">© 2025 MONCAR RMH — Todos los derechos reservados</div>
</footer>

<!-- FORM PEDIDO -->
<div class="fmodal-ov" id="fmodal-ov">
  <div class="fmodal">
    <div class="fmodal-title">Datos para tu pedido</div>
    <div class="fmodal-sub">Necesitamos estos datos para asignarte el representante de tu zona y coordinar la entrega.</div>
    <div class="fg">
      <label>Nombre *</label>
      <input type="text" id="f-nombre" placeholder="Tu nombre completo">
    </div>
    <div class="fg">
      <label>Nombre del negocio</label>
      <input type="text" id="f-negocio" placeholder="Ferreteria, taller, bulonera, local... (opcional)">
      <div class="fg-note">Opcional si sos particular</div>
    </div>
    <div class="fg">
      <label>Localidad / Ciudad o Codigo postal *</label>
      <input type="text" id="f-localidad" placeholder="Ej: Mendoza, Mar del Plata, 1900..." oninput="detectarZonaInput()">
      <div class="zona-res" id="zona-res"></div>
    </div>
    <div class="fmodal-acts">
      <button class="btn-cnc" onclick="cerrarFormModal()">Cancelar</button>
      <button class="btn-cfm" onclick="confirmarPedido()">Enviar pedido &#8594;</button>
    </div>
  </div>
</div>

<!-- CART -->
<div class="cart-overlay" id="overlay" onclick="cerrarCart()"></div>
<div class="cart-panel" id="cart-panel">
  <div class="cart-header"><h2>Tu pedido</h2><button class="btn-xc" onclick="cerrarCart()">&#x2715;</button></div>
  <div class="cart-items" id="cart-items"></div>
  <div class="cart-footer">
    <div class="cart-total"><span>Total estimado</span><span id="cart-total">$0</span></div>
    <button class="btn-wa-c" onclick="pedirPorWA()">&#128172; Pedir por WhatsApp</button>
    <button class="btn-em-c" onclick="pedirPorEmail()">Pedir por Email</button>
  </div>
</div>

<script>
const PRODUCTOS = {productos_json};

const SVG_MC=`<svg width="50" height="50" viewBox="0 0 50 50" fill="none"><circle cx="25" cy="25" r="17" fill="none" stroke="#FF6B00" stroke-width="1.8" opacity=".2"/><circle cx="25" cy="25" r="9" fill="none" stroke="#FF6B00" stroke-width="1.8" opacity=".45"/><circle cx="25" cy="25" r="3" fill="#FF6B00" opacity=".8"/><line x1="9" y1="25" x2="16" y2="25" stroke="#FF6B00" stroke-width="1.8" stroke-linecap="round" opacity=".4"/><line x1="34" y1="25" x2="41" y2="25" stroke="#FF6B00" stroke-width="1.8" stroke-linecap="round" opacity=".4"/><line x1="25" y1="9" x2="25" y2="16" stroke="#FF6B00" stroke-width="1.8" stroke-linecap="round" opacity=".4"/><line x1="25" y1="34" x2="25" y2="41" stroke="#FF6B00" stroke-width="1.8" stroke-linecap="round" opacity=".4"/></svg>`;
const SVG_RMH=`<svg width="50" height="50" viewBox="0 0 50 50" fill="none"><rect x="9" y="15" width="32" height="20" rx="4" stroke="#FF6B00" stroke-width="1.8" opacity=".3" fill="none"/><rect x="15" y="20" width="20" height="10" rx="2" fill="#FF6B00" opacity=".1"/><circle cx="25" cy="25" r="3" fill="#FF6B00" opacity=".7"/><line x1="18" y1="25" x2="22" y2="25" stroke="#FF6B00" stroke-width="1.8" stroke-linecap="round" opacity=".5"/><line x1="28" y1="25" x2="32" y2="25" stroke="#FF6B00" stroke-width="1.8" stroke-linecap="round" opacity=".5"/></svg>`;

let cart = JSON.parse(localStorage.getItem('moncar_cart')||'[]');
let filtroLinea='';

function fmt(n){{ return n ? '$'+Math.round(n).toLocaleString('es-AR') : 'Consultar'; }}

function setNavActive(id){{
  ['inicio','productos','nosotros'].forEach(n=>document.getElementById('nav-'+n).classList.toggle('active',n===id));
}}
function irAInicio(){{
  document.querySelectorAll('.page').forEach(p=>p.classList.remove('active'));
  document.getElementById('page-inicio').classList.add('active');
  setNavActive('inicio'); window.scrollTo(0,0);
}}
function irAProductos(cat,marca){{
  document.querySelectorAll('.page').forEach(p=>p.classList.remove('active'));
  document.getElementById('page-productos').classList.add('active');
  setNavActive('productos');
  // Reset sidebar checkboxes
  document.querySelectorAll('#sb-cats input,#sb-marcas input').forEach(cb=>cb.checked=false);
  if(cat){{
    const cb=document.querySelector(`#sb-cats input[value="${{cat}}"]`);
    if(cb) cb.checked=true;
  }}
  if(marca){{
    const cb=document.querySelector(`#sb-marcas input[value="${{marca}}"]`);
    if(cb) cb.checked=true;
  }}
  filtrar(); window.scrollTo(0,0);
}}
function irANosotros(){{
  document.querySelectorAll('.page').forEach(p=>p.classList.remove('active'));
  document.getElementById('page-nosotros').classList.add('active');
  setNavActive('nosotros'); window.scrollTo(0,0);
}}
function irACat(c){{ irAProductos(c,null); }}
function irAMarca(m){{ irAProductos(null,m); }}

function setLinea(btn){{
  document.querySelectorAll('.sb-linea-btn').forEach(b=>b.classList.remove('active'));
  btn.classList.add('active');
  filtroLinea=btn.dataset.val;
  filtrar();
}}

function filtrar(){{
  const q=document.getElementById('search').value.toLowerCase().trim();
  const soloStock=document.getElementById('solo-stock').checked;
  document.getElementById('clear-btn').style.display=q?'block':'none';

  const selCats=[...document.querySelectorAll('#sb-cats input:checked')].map(c=>c.value);
  const selMarcas=[...document.querySelectorAll('#sb-marcas input:checked')].map(c=>c.value);
  const pMin=parseFloat(document.getElementById('precio-min').value)||0;
  const pMax=parseFloat(document.getElementById('precio-max').value)||Infinity;
  const filEsp=document.getElementById('fil-esp').value;
  const filAnch=document.getElementById('fil-anch').value;
  const filLarg=document.getElementById('fil-larg').value;
  const filTerm=document.getElementById('fil-term').value;

  // Active tags
  const af=document.getElementById('active-filters');
  const tags=[];
  if(filtroLinea) tags.push(`<div class="active-tag" onclick="document.querySelector('.sb-linea-btn[data-val=\\'\\']').click()">Linea: ${{filtroLinea}} ✕</div>`);
  selCats.forEach(c=>tags.push(`<div class="active-tag" onclick="uncheckCat('${{c}}')">📂 ${{c}} ✕</div>`));
  selMarcas.forEach(m=>tags.push(`<div class="active-tag" onclick="uncheckMarca('${{m}}')">🏷️ ${{m}} ✕</div>`));
  if(pMin>0||pMax<Infinity) tags.push(`<div class="active-tag" onclick="limpiarPrecio()">Precio ✕</div>`);
  af.innerHTML=tags.join('');
  af.style.display=tags.length?'flex':'none';

  const visibles=PRODUCTOS.filter(p=>{{
    if(filtroLinea&&p.linea!==filtroLinea) return false;
    if(soloStock&&!p.stock) return false;
    if(selCats.length&&!selCats.includes(p.cat)) return false;
    if(selMarcas.length&&!selMarcas.includes(p.marca)) return false;
    if(p.precio>0&&(p.precio<pMin||p.precio>pMax)) return false;
    if(filEsp&&String(p.esp)!==filEsp) return false;
    if(filAnch&&String(p.anch)!==filAnch) return false;
    if(filLarg&&String(p.larg)!==filLarg) return false;
    if(filTerm&&String(p.terminal)!==filTerm) return false;
    if(q&&!(p.cod+' '+p.desc+' '+p.marca+' '+p.cat).toLowerCase().includes(q)) return false;
    return true;
  }});

  renderGrid(visibles);
  document.getElementById('results-info').textContent=visibles.length+' producto'+(visibles.length!==1?'s':'');
  document.getElementById('no-results').style.display=visibles.length?'none':'block';
  document.getElementById('grid').style.display=visibles.length?'grid':'none';
}}

function renderGrid(lista){{
  document.getElementById('grid').innerHTML=lista.map(p=>{{
    const svg=p.linea==='RMH'?SVG_RMH:SVG_MC;
    const enCart=cart.find(c=>c.cod===p.cod);
    const btnTxt=p.stock?(enCart?'Agregar mas':'Agregar'):'Sin stock';
    const marcaTxt=(p.marca&&!['RMH','Moncar','Electrodomesticos'].includes(p.marca))?p.marca:'';
    return `<div class="card">
      <div class="card-icon">${{svg}}</div>
      <div class="card-body">
        <div class="card-badges"><span class="badge bl">${{p.linea}}</span><span class="badge bc">${{p.cat}}</span></div>
        <div class="card-cod">${{p.cod}}</div>
        <div class="card-desc">${{p.desc}}</div>
        ${{marcaTxt?`<div class="card-sub">${{marcaTxt}}</div>`:''}}
        ${{p.dims?`<div class="card-sub">${{p.dims}} mm</div>`:''}}
        <div class="card-precio">${{fmt(p.precio)}}</div>
        <span class="card-stock ${{p.stock?'sok':'sno'}}">${{p.stock?'En stock':'Sin stock'}}</span>
        <button class="btn-add" ${{p.stock?'':'disabled'}} onclick="agregarAlCarrito('${{p.cod}}')">${{btnTxt}}</button>
      </div></div>`;
  }}).join('');
}}

function limpiarBusqueda(){{ document.getElementById('search').value=''; filtrar(); }}
function limpiarPrecio(){{ document.getElementById('precio-min').value=''; document.getElementById('precio-max').value=''; filtrar(); }}
function uncheckCat(v){{ const cb=document.querySelector(`#sb-cats input[value="${{v}}"]`); if(cb)cb.checked=false; filtrar(); }}
function uncheckMarca(v){{ const cb=document.querySelector(`#sb-marcas input[value="${{v}}"]`); if(cb)cb.checked=false; filtrar(); }}
function limpiarFiltros(){{
  document.querySelectorAll('#sb-cats input,#sb-marcas input').forEach(cb=>cb.checked=false);
  document.getElementById('solo-stock').checked=false;
  document.getElementById('precio-min').value='';
  document.getElementById('precio-max').value='';
  document.getElementById('fil-esp').value='';
  document.getElementById('fil-anch').value='';
  document.getElementById('fil-larg').value='';
  document.getElementById('fil-term').value='';
  document.getElementById('search').value='';
  document.querySelector('.sb-linea-btn[data-val=""]').click();
}}

function abrirSidebar(){{ document.getElementById('sidebar').classList.add('open'); document.getElementById('sb-overlay').classList.add('open'); }}
function cerrarSidebar(){{ document.getElementById('sidebar').classList.remove('open'); document.getElementById('sb-overlay').classList.remove('open'); }}

// CART
function agregarAlCarrito(cod){{
  const p=PRODUCTOS.find(x=>x.cod===cod); if(!p||!p.stock) return;
  const ex=cart.find(c=>c.cod===cod);
  if(ex) ex.qty++; else cart.push({{cod:p.cod,desc:p.desc,precio:p.precio,qty:1}});
  guardarCart(); filtrar(); abrirCart();
}}
function cambiarQty(cod,delta){{
  const ex=cart.find(c=>c.cod===cod); if(!ex) return;
  ex.qty+=delta; if(ex.qty<=0) cart=cart.filter(c=>c.cod!==cod);
  guardarCart();
}}
function quitarItem(cod){{ cart=cart.filter(c=>c.cod!==cod); guardarCart(); }}
function guardarCart(){{
  localStorage.setItem('moncar_cart',JSON.stringify(cart));
  renderCart();
  document.getElementById('cart-count').textContent=cart.reduce((s,c)=>s+c.qty,0);
}}
function renderCart(){{
  const el=document.getElementById('cart-items');
  if(!cart.length){{ el.innerHTML='<div class="cart-empty">El carrito esta vacio</div>'; document.getElementById('cart-total').textContent='$0'; return; }}
  el.innerHTML=cart.map(c=>`<div class="ci">
    <div class="ci-info">
      <div class="ci-cod">${{c.cod}}</div>
      <div class="ci-desc">${{c.desc}}</div>
      <div class="ci-precio">${{fmt(c.precio)}}</div>
      <div class="ci-qty">
        <button onclick="cambiarQty('${{c.cod}}',-1);renderCart();filtrar()">-</button>
        <span>${{c.qty}}</span>
        <button onclick="cambiarQty('${{c.cod}}',1);renderCart();filtrar()">+</button>
      </div>
    </div>
    <button class="btn-rm" onclick="quitarItem('${{c.cod}}');renderCart();filtrar()">✕</button>
  </div>`).join('');
  document.getElementById('cart-total').textContent=fmt(cart.reduce((s,c)=>s+c.precio*c.qty,0));
}}
function abrirCart(){{ document.getElementById('cart-panel').classList.add('open'); document.getElementById('overlay').style.display='block'; renderCart(); }}
function cerrarCart(){{ document.getElementById('cart-panel').classList.remove('open'); document.getElementById('overlay').style.display='none'; }}
// ---- ZONAS Y ROUTING ----
const HERMANOS = [
  {{nombre:'Mapi', wa:'5491122334455', email:'mapi@moncarrmh.com',
    zonas:['MENDOZA','NORTE DE CORDOBA','SAN JUAN','SAN LUIS','RIO NEGRO',
           'VILLA LA ANGOSTURA','CHOS MALAL','RAWSON','TRELEW',
           'COMODORO RIVADAVIA','PUERTO MADRYN','SANTA CRUZ','NEUQUEN']}},
  {{nombre:'Juli', wa:'5491133445566', email:'juli@moncarrmh.com',
    zonas:['SANTIAGO DEL ESTERO','TUCUMAN','LUJAN','PERGAMINO','VENADO TUERTO',
           'LA CARLOTA','RIO CUARTO','SAMPACHO','VILLA MERCEDES',
           'CANADA DE GOMEZ','ARMSTRONG','MARCOS JUAREZ','MARCO JUAREZ',
           'LEONES','BELL VILLE','VILLA MARIA','SANTA FE NORTE',
           'LA RIOJA','CATAMARCA']}},
  {{nombre:'Luvi', wa:'5491144556677', email:'luvi@moncarrmh.com',
    zonas:['ENTRE RIOS','CORRIENTES','MISIONES','SANTA FE',
           'ZARATE','CAPILLA','JOSE C. PAZ','JOSE C PAZ']}},
  {{nombre:'Fede', wa:'5491155667788', email:'fede@moncarrmh.com',
    zonas:['BAHIA BLANCA','SUIPACHA','PELLEGRINI','CHACABUCO','VEDIA',
           'ROQUE PEREZ','LAS FLORES','AYACUCHO','LA PAMPA',
           'CORDOBA SUR','TRES ARROYOS','MAR DEL PLATA','ZAPALA','GRUTAS']}}
];
const FALLBACK = {{nombre:null, wa:'5491199887766', email:'contacto@moncarrmh.com'}};

function norm(s){{
  return s.toUpperCase().normalize('NFD').replace(/[\\u0300-\\u036f]/g,'').trim();
}}
function detectarZona(localidad){{
  const loc=norm(localidad);
  if(loc.length<2) return [];
  let matches=[];
  for(const h of HERMANOS){{
    let best=0, bz='';
    for(const z of h.zonas){{
      const zn=norm(z); if(zn.length<3) continue;
      let s=0;
      if(loc===zn) s=1000;
      else if(loc.includes(zn)) s=zn.length*10;
      else if(zn.includes(loc)&&loc.length>=4) s=loc.length;
      if(s>best){{best=s;bz=z;}}
    }}
    if(best>0) matches.push({{h,score:best,zona:bz}});
  }}
  return matches.sort((a,b)=>b.score-a.score);
}}

let _pedidoTipo='';
function pedirPorWA(){{ if(!cart.length) return; _pedidoTipo='wa'; abrirFormModal(); }}
function pedirPorEmail(){{ if(!cart.length) return; _pedidoTipo='email'; abrirFormModal(); }}

function abrirFormModal(){{
  document.getElementById('fmodal-ov').classList.add('open');
  document.getElementById('zona-res').style.display='none';
  ['f-nombre','f-localidad'].forEach(id=>document.getElementById(id).classList.remove('ferr'));
}}
function cerrarFormModal(){{
  document.getElementById('fmodal-ov').classList.remove('open');
}}
function detectarZonaInput(){{
  const loc=document.getElementById('f-localidad').value.trim();
  const el=document.getElementById('zona-res');
  if(loc.length<3){{el.style.display='none';return;}}
  const m=detectarZona(loc);
  el.style.display='block';
  el.className='zona-res';
  if(!m.length){{
    el.classList.add('zwarn');
    el.innerHTML='&#9888;&#65039; No reconocemos esa localidad. Igual podes enviar el pedido y te contactamos.';
  }} else if(m.length===1||m[0].score>=(m[1]&&m[1].score||0)*1.5){{
    el.classList.add('zok');
    el.innerHTML='&#10003; Tu zona corresponde a <strong>'+m[0].h.nombre+'</strong>, quien se va a contactar con vos.';
  }} else {{
    el.classList.add('zwarn');
    el.innerHTML='&#9888;&#65039; Puede corresponder a <strong>'+m.slice(0,2).map(x=>x.h.nombre).join(' o ')+'</strong>. Especifica un poco mas la localidad.';
  }}
}}
function confirmarPedido(){{
  const nombre=document.getElementById('f-nombre').value.trim();
  const negocio=document.getElementById('f-negocio').value.trim();
  const localidad=document.getElementById('f-localidad').value.trim();
  let ok=true;
  document.getElementById('f-nombre').classList.toggle('ferr',!nombre); if(!nombre) ok=false;
  document.getElementById('f-localidad').classList.toggle('ferr',!localidad); if(!localidad) ok=false;
  if(!ok) return;

  const m=detectarZona(localidad);
  let dest=FALLBACK;
  if(m.length===1||(m.length>1&&m[0].score>=(m[1]&&m[1].score||0)*1.5)) dest=m[0].h;

  const linea1='Hola'+(dest.nombre?' '+dest.nombre:'')+'! Nuevo pedido desde la web:';
  let msg=linea1+'\\n\\n';
  msg+='Cliente: '+nombre+'\\n';
  if(negocio) msg+='Negocio: '+negocio+'\\n';
  msg+='Localidad: '+localidad+'\\n\\n';
  msg+='Productos:\\n';
  cart.forEach(c=>{{msg+=`- ${{c.cod}} x${{c.qty}}\\n`;}});
  msg+='\\nMuchas gracias!';

  cerrarFormModal();

  if(_pedidoTipo==='wa'){{
    window.open('https://wa.me/'+dest.wa+'?text='+encodeURIComponent(msg),'_blank');
  }} else {{
    const subj='Nuevo pedido MONCAR RMH'+(negocio?' - '+negocio:'');
    window.location.href='mailto:'+dest.email+'?subject='+encodeURIComponent(subj)+'&body='+encodeURIComponent(msg);
  }}
}}

guardarCart(); filtrar();
</script>
</body>
</html>'''

# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
if __name__ == '__main__':
    print("Leyendo stock MONCAR...")
    stock_mc = leer_stock_moncar()

    print("Leyendo productos MONCAR desde LP...")
    productos_mc = leer_moncar_lp(stock_mc)
    print(f"  {len(productos_mc)} productos MONCAR")

    print("Leyendo productos RMH...")
    productos_rmh = leer_rmh()
    print(f"  {len(productos_rmh)} productos RMH")

    todos = productos_mc + productos_rmh

    print("\n=== Distribucion por categoria ===")
    cats = Counter(p['cat'] for p in todos)
    for cat, n in sorted(cats.items(), key=lambda x: -x[1]):
        print(f"  {cat}: {n}")

    EXCLUIR = {'Moncar','Electrodomesticos','Cobreados','Porta Carbones','Power Tools',
               'Experto','Bta','Neo','Lazo Asincron','Versa','Lavarropas','Aspiradora',
               'Ventilador','Batidora','Maquina De Coser','Lustra Aspiradora','Hidrolavadora',
               'Enceradoras','Afeitadoras','Bordeadora','Barre Alfombras','Salkor',
               'Stingray','Dover','Gladiator','Stayer','Peugeot','Dowen Paggio',
               'Martins','Agro','Argentec','Adiabatic'}
    mc = Counter(p['marca'] for p in todos if p['linea']=='MONCAR' and p['marca'] not in EXCLUIR)
    destacadas = [(m,c) for m,c in mc.items() if m in MARCAS_CONFIG]
    print(f"\n=== Marcas destacadas para homepage ({len(destacadas)}) ===")
    for m,c in sorted(destacadas, key=lambda x: -x[1]):
        print(f"  {m}: {c}")

    print("\nGenerando HTML...")
    html = generar_html(todos)
    with open(OUT_PATH, 'w', encoding='utf-8') as f:
        f.write(html)

    print(f"\nSitio generado: {OUT_PATH} ({len(html)//1024} KB)")
